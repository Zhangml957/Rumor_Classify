from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORT_DOCX = ROOT / "report.docx"
ASSET_DIR = ROOT / "outputs" / "report_assets"


def set_page(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)


def set_font(run, name: str = "宋体", size: float = 12, bold: bool = False, color: str | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, first_line_chars: float | None = None, space_after: float = 6) -> None:
    fmt = paragraph.paragraph_format
    fmt.line_spacing = 1.25
    fmt.space_after = Pt(space_after)
    if first_line_chars is not None:
        fmt.first_line_indent = Pt(24 * first_line_chars)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_font(run, name="黑体", size=14 if level == 1 else 12, bold=True, color="1F4E79")
    p.paragraph_format.space_before = Pt(8 if level == 1 else 4)
    p.paragraph_format.space_after = Pt(4)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_font(run, name="宋体", size=11)
    style_paragraph(p, first_line_chars=2.0, space_after=4)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    set_font(run, name="宋体", size=11)
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_after = Pt(2)


def set_cell_text(cell, text: str, bold: bool = False, size: float = 10.5) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    set_font(run, name="宋体", size=size, bold=bold)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(0)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_result_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.autofit = False
    widths = [Inches(2.8), Inches(1.5), Inches(2.2)]
    headers = ["方案", "准确率", "说明"]
    for cell, text, width in zip(table.rows[0].cells, headers, widths):
        cell.width = width
        set_cell_text(cell, text, bold=True)
        shade_cell(cell, "D9E7F5")
    rows = [
        ("Hybrid Retrieval", "84.29%", "仅使用检索投票"),
        ("BERTweet + Hybrid", "85.54%", "Transformer 主分类 + 检索"),
        ("BERTweet + Rules", "89.28%", "加入证据冲突与官方更新纠偏"),
        ("Event Prefix", "87.03%", "显式加入事件前缀，效果下降"),
        ("LLM Semantic", "86.03%", "真实 LLM 语义改判，整体下降"),
        ("Final Mainline", "90.52%", "加入 rumor amplification 与安全通知纠偏"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].width = widths[i]
            set_cell_text(cells[i], value)


def add_example_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.autofit = False
    widths = [Inches(1.4), Inches(2.2), Inches(2.9)]
    headers = ["示例类型", "输入文本片段", "中文判断依据示例"]
    for cell, text, width in zip(table.rows[0].cells, headers, widths):
        cell.width = width
        set_cell_text(cell, text, bold=True)
        shade_cell(cell, "E8EEF5")
    rows = [
        (
            "安全通知",
            "UPDATE: #uOttawa courses + exams officially cancelled... Lockdown still in effect. Stay indoors, stay safe...",
            "该推文被判定为非谣言，因为它在传播中更像一条机构安全通知。“officially cancelled”“lockdown still in effect”“stay indoors”等表达体现了事务安排和安全指令属性，而不是未证实指控的传播。",
        ),
        (
            "谣言式传播",
            "Shoot unarmed kid. Conceal evidence. Impose martial law. Smear the victim...",
            "该推文被判定为谣言，因为它在传播中更像一条对未证实指控的谣言式传播。“conceal evidence”“smear the victim”“martial law”等表述带有强指控和放大传播特征，检索到的相似样本也支持这一判断。",
        ),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].width = widths[i]
            set_cell_text(cells[i], value, size=10)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_font(run, name="宋体", size=10.5, color="555555")
    p.paragraph_format.space_after = Pt(6)


def build_report() -> None:
    doc = Document()
    set_page(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("2026《人工智能导论》大作业报告")
    set_font(run, name="黑体", size=18, bold=True, color="1F4E79")
    title.paragraph_format.space_after = Pt(8)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("任务名称：可解释的谣言检测")
    set_font(run, name="宋体", size=12, bold=True)
    subtitle.paragraph_format.space_after = Pt(10)

    meta_lines = [
        "完成组号：待填写",
        "小组人员：待填写",
        "完成时间：2026年6月4日",
        "代码仓库：https://github.com/SiriThree/Rumor_Classify",
    ]
    for line in meta_lines:
        p = doc.add_paragraph()
        run = p.add_run(line)
        set_font(run, name="宋体", size=11)
        p.paragraph_format.space_after = Pt(2)

    add_heading(doc, "1．任务目标")
    add_body(
        doc,
        "本项目面向英文社交媒体推文的谣言检测任务，输入一条推文后输出二分类结果与相应判断依据。我们以 val.csv 上的准确率和可解释性为核心目标，构建了“Transformer 主分类 + 检索证据 + 规则纠偏 + 解释生成”的复合系统，在保证可复现的前提下尽量提升检测效果。"
    )

    add_heading(doc, "2．具体内容")
    add_heading(doc, "（1）实施方案", level=2)
    add_body(
        doc,
        "最终方案采用 BERTweet 作为主分类器，对 tweet 进行二分类微调；同时使用 lexical retrieval 与 dense retrieval 从训练集检索相似样本，并通过 RRF 进行混合排序。对于高风险样本，系统再结合证据一致性、官方通知模式、rumor amplification 模式和安全通知模式进行纠偏，最后输出预测标签和解释文本。"
    )
    add_bullet(doc, "主分类器：vinai/bertweet-base，适合 tweet 语域。")
    add_bullet(doc, "检索证据：稀疏检索 + 句向量检索。")
    add_bullet(doc, "融合策略：强一致覆盖、证据冲突复判、官方更新修正、rumor amplification 修正、安全通知修正。")
    add_bullet(doc, "解释模块：基于文本信号、相似证据和最终标签生成中文判断依据。")

    add_heading(doc, "（2）核心代码分析", level=2)
    add_body(
        doc,
        "系统主流程集中在 src/rumor_system/pipeline.py。dataset.py 负责数据读入与规范化；transformer.py 负责 BERTweet 的训练、保存和推理；hybrid.py 将 lexical 与 dense 两路证据进行融合；pipeline.py 中的 fuse_prediction 实现了检索增强决策逻辑，是本项目提分的关键模块。"
    )
    add_bullet(doc, "数据层：读取 id、text、label、event，并生成 normalized_text。")
    add_bullet(doc, "模型层：对 BERTweet 进行监督微调，输出 label 与 confidence。")
    add_bullet(doc, "检索层：从 train.csv 中找相似推文，为预测提供证据支持。")
    add_bullet(doc, "决策层：围绕高频错例做窄规则修正，而不是全局硬覆盖。")

    add_heading(doc, "（3）检测结果分析（正确率等）", level=2)
    add_body(
        doc,
        "最终稳定主线在 val.csv 上达到 363/401，准确率 90.52%。从实验过程看，真正有效的提升主要来自“检索证据 + 面向错例的定向纠偏”，而不是盲目换 backbone 或直接引入 LLM 改判。我们还验证了事件前缀输入和 LLM semantic analyzer，两者在当前数据上都没有带来净收益。"
    )
    add_result_table(doc)

    if (ASSET_DIR / "accuracy_compare.png").exists():
        doc.add_picture(str(ASSET_DIR / "accuracy_compare.png"), width=Inches(5.9))
        add_caption(doc, "图1 不同方案在验证集上的准确率对比")

    add_body(
        doc,
        "从事件维度看，Germanwings 与 SydneySiege 最难，准确率分别为 80.43% 和 88.43%；Ferguson 为 90.83%，Ottawa 为 94.38%。这说明数据中的错误并非随机分布，而是高度依赖事件语境与表达模式。"
    )
    if (ASSET_DIR / "event_accuracy.png").exists():
        doc.add_picture(str(ASSET_DIR / "event_accuracy.png"), width=Inches(5.9))
        add_caption(doc, "图2 各事件子集上的准确率")

    add_body(
        doc,
        "需要说明的是，训练集与验证集覆盖相同事件集合，因此当前高准确率部分受益于事件内相似表达、近重复模板和检索增强能力。也就是说，该结果在当前验证集上有效，但对全新事件的泛化能力仍需进一步验证。"
    )

    add_heading(doc, "（4）判断依据的分析（可解释性等）", level=2)
    add_body(
        doc,
        "解释模块并不让大模型直接决定标签，而是基于最终预测、相似训练样本、top evidence 标签和置信度生成中文判断依据。新版 explanation 采用“三段式”结构：先说明当前 tweet 在传播中扮演的角色，再指出触发判断的关键短语，最后补充检索证据如何支持或限制该结论。这样既能提高可读性，也能让判断依据更加贴近真实决策过程。"
    )
    add_bullet(doc, "对于 rumor amplification 类文本，系统更关注“未证实指控的传播作用”，而不是表面情绪强弱。")
    add_bullet(doc, "对于官方公告、安全通知、课程取消等文本，系统更关注“通知角色”，而不是事件背景中的 shooting/lockdown 词汇。")
    add_bullet(doc, "对于 headline-style 新闻快讯，系统会优先结合 top evidence 与事件风格判断其是否真正属于 rumor thread。")
    add_example_table(doc)
    add_caption(doc, "表1 中文判断依据示例")

    add_heading(doc, "3．工作总结")
    add_heading(doc, "（1）收获、心得", level=2)
    add_body(
        doc,
        "本次大作业最大的收获是认识到：rumor detection 并不等同于真假判断。很多文本是否属于 rumor，更取决于它在事件早期传播中的角色。单纯依赖表层词汇或情绪强度容易产生伪相关，只有把深度学习、检索证据和错误分析结合起来，系统效果才会稳定提升。"
    )

    add_heading(doc, "（2）遇到问题及解决思路", level=2)
    add_body(
        doc,
        "项目中遇到的主要问题包括：headline false alarm、opinion-like rumor miss、检索证据带偏，以及在线接口推理较慢。对应地，我们通过错误样本分型、增加缓存、放弃副作用较大的 LLM 改判、保留有效的窄规则纠偏，最终将主线准确率提升到了 90.52%。"
    )

    add_heading(doc, "4．课程建议")
    add_body(
        doc,
        "如果后续课程继续使用类似任务，建议在数据中提供 conversation thread、source tweet、reply 关系、发布时间或来源类型等上下文信息。对于 PHEME 风格任务，单 tweet 分类天然受限；若能加入 thread 结构和时间线信息，模型的真实性能与可解释性都还有明显提升空间。"
    )

    doc.save(str(REPORT_DOCX))


if __name__ == "__main__":
    build_report()
